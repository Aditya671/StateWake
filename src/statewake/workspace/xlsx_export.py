"""Narrow openpyxl adapter for human-friendly StateWake XLSX exports."""

from __future__ import annotations

import hashlib
import json
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from .errors import WorkspaceRepositoryError
from .projections import NormalizedRecordProjection

XLSX_FORMAT = "xlsx"
XLSX_SCHEMA_VERSION = "1"
XLSX_RECORD_SHEET = "Records"
XLSX_METADATA_SHEET = "Metadata"
XLSX_RECORD_COLUMNS = (
    "record_id",
    "receipt_id",
    "artifact_digest",
    "artifact_size",
    "producer_id",
    "producer_type",
    "producer_version",
    "source_ref",
    "source_event_id",
    "run_id",
    "captured_at",
    "sensitivity",
    "policy_id",
    "verification_status",
    "reliability_state",
    "created_at",
)
XLSX_METADATA_COLUMNS = ("key", "value")


class XlsxExportError(WorkspaceRepositoryError):
    """Raised when an XLSX export cannot be created or verified."""


@dataclass(frozen=True, slots=True)
class XlsxExportResult:
    """Describe one successfully written XLSX export."""

    export_id: str
    output_path: Path
    output_digest: str
    row_count: int
    schema_version: str
    disclosure_max_sensitivity: str


def write_xlsx(
    records: tuple[NormalizedRecordProjection, ...],
    output: Path,
    *,
    schema_version: str,
    disclosure_max_sensitivity: str,
    query_definition: Mapping[str, object],
) -> XlsxExportResult:
    """Write one bounded projection using a narrow optional Excel adapter."""
    openpyxl = _load_openpyxl()
    output = Path(output).expanduser()
    if output.exists() and output.is_dir():
        raise XlsxExportError("XLSX output requires a file path.")
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    records_sheet = workbook.active
    records_sheet.title = XLSX_RECORD_SHEET
    records_sheet.freeze_panes = "A2"
    records_sheet.auto_filter.ref = _data_range(len(records), len(XLSX_RECORD_COLUMNS))
    records_sheet.append(XLSX_RECORD_COLUMNS)
    for cell in records_sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for record in records:
        records_sheet.append(_record_row(record))
        _force_string_formula_safe(records_sheet[records_sheet.max_row], openpyxl)
    _set_column_widths(records_sheet, openpyxl)

    metadata_sheet = workbook.create_sheet(XLSX_METADATA_SHEET)
    metadata_sheet.freeze_panes = "A2"
    metadata_sheet.append(XLSX_METADATA_COLUMNS)
    for cell in metadata_sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    metadata = (
        ("format", XLSX_FORMAT),
        ("schema_version", schema_version),
        ("row_count", len(records)),
        ("disclosure_max_sensitivity", disclosure_max_sensitivity),
        (
            "query_definition",
            json.dumps(
                dict(query_definition),
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
        ),
    )
    for key, value in metadata:
        metadata_sheet.append((key, value))
    _set_metadata_widths(metadata_sheet, openpyxl)

    properties = workbook.properties
    properties.creator = "StateWake"
    properties.lastModifiedBy = "StateWake"
    properties.created = datetime(2000, 1, 1, tzinfo=UTC)
    properties.modified = datetime(2000, 1, 1, tzinfo=UTC)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    workbook.calculation.calcMode = "auto"

    temporary_path = output.with_name(f".{output.name}.{uuid4().hex}.tmp")
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output)
    except (OSError, ValueError) as exc:
        try:
            temporary_path.unlink()
        except OSError:
            pass
        raise XlsxExportError(f"XLSX export could not be written: {exc}") from exc

    return XlsxExportResult(
        export_id=str(uuid4()),
        output_path=output,
        output_digest=_sha256_file(output),
        row_count=len(records),
        schema_version=schema_version,
        disclosure_max_sensitivity=disclosure_max_sensitivity,
    )


def verify_xlsx(
    path: Path,
    *,
    expected_rows: int,
    expected_schema_version: str,
) -> None:
    """Reopen XLSX and validate sheets, columns, metadata, and row count."""
    openpyxl = _load_openpyxl()
    path = Path(path)
    if not path.is_file():
        raise XlsxExportError(f"XLSX export does not exist: {path}")
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        if tuple(workbook.sheetnames) != (XLSX_RECORD_SHEET, XLSX_METADATA_SHEET):
            raise XlsxExportError(
                "XLSX sheet mismatch: "
                f"expected {(XLSX_RECORD_SHEET, XLSX_METADATA_SHEET)!r}, "
                f"got {tuple(workbook.sheetnames)!r}."
            )
        records_sheet = workbook[XLSX_RECORD_SHEET]
        header = tuple(next(records_sheet.iter_rows(values_only=True)))
        if header != XLSX_RECORD_COLUMNS:
            raise XlsxExportError(
                "XLSX schema mismatch: "
                f"expected {XLSX_RECORD_COLUMNS!r}, got {header!r}."
            )
        row_count = sum(1 for _ in records_sheet.iter_rows(min_row=2, values_only=True))
        if row_count != expected_rows:
            raise XlsxExportError(
                f"XLSX row count mismatch: expected {expected_rows}, got {row_count}."
            )
        metadata_sheet = workbook[XLSX_METADATA_SHEET]
        metadata = {
            str(row[0]): row[1]
            for row in metadata_sheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None
        }
        if metadata.get("schema_version") != expected_schema_version:
            raise XlsxExportError(
                "XLSX metadata schema mismatch: "
                f"expected {expected_schema_version!r}, "
                f"got {metadata.get('schema_version')!r}."
            )
        if metadata.get("format") != XLSX_FORMAT:
            raise XlsxExportError(
                f"XLSX metadata format mismatch: expected {XLSX_FORMAT!r}, "
                f"got {metadata.get('format')!r}."
            )
    except XlsxExportError:
        raise
    except (OSError, ValueError, KeyError, TypeError) as exc:
        raise XlsxExportError(f"XLSX export could not be reopened: {exc}") from exc
    finally:
        try:
            workbook.close()
        except UnboundLocalError:
            pass


def _record_row(record: NormalizedRecordProjection) -> tuple[object, ...]:
    """Convert one normalized record to stable Excel cell values."""
    return (
        record.record_id,
        record.receipt_id,
        record.artifact_digest,
        record.artifact_size,
        record.producer_id,
        record.producer_type,
        record.producer_version,
        record.source_ref,
        record.source_event_id,
        record.run_id,
        record.captured_at.astimezone(UTC).replace(tzinfo=None),
        record.sensitivity,
        record.policy_id,
        record.verification_status,
        record.reliability_state,
        record.created_at.astimezone(UTC).replace(tzinfo=None),
    )


def _force_string_formula_safe(row: Any, openpyxl: Any) -> None:
    """Prevent formula interpretation without changing exported string values."""
    for cell in row:
        value = cell.value
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            cell.data_type = "s"


def _set_column_widths(sheet: Any, openpyxl: Any) -> None:
    """Apply bounded human-readable widths to the record sheet."""
    widths = (22, 22, 68, 14, 22, 18, 18, 32, 22, 22, 24, 16, 18, 22, 22, 24)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width


def _set_metadata_widths(sheet: Any, openpyxl: Any) -> None:
    """Apply readable widths to the metadata sheet."""
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 110


def _data_range(row_count: int, column_count: int) -> str:
    """Return the bounded autofilter range for a record sheet."""
    end_column = _column_letter(column_count)
    end_row = max(row_count + 1, 1)
    return f"A1:{end_column}{end_row}"


def _column_letter(number: int) -> str:
    """Convert one positive one-based column number to an Excel letter."""
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def _load_openpyxl() -> Any:
    """Import openpyxl lazily so SQLite and non-XLSX paths stay lightweight."""
    try:
        import openpyxl
    except ImportError as exc:
        raise XlsxExportError(
            "XLSX export requires the optional openpyxl dependency."
        ) from exc
    return openpyxl


def _sha256_file(path: Path) -> str:
    """Return the SHA-256 identity of one XLSX output file."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


__all__ = [
    "XLSX_FORMAT",
    "XLSX_METADATA_COLUMNS",
    "XLSX_METADATA_SHEET",
    "XLSX_RECORD_COLUMNS",
    "XLSX_RECORD_SHEET",
    "XLSX_SCHEMA_VERSION",
    "XlsxExportError",
    "XlsxExportResult",
    "verify_xlsx",
    "write_xlsx",
]
