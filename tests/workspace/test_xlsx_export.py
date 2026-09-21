"""Regression tests for the Tier 9 XLSX export surface."""

from __future__ import annotations

import sqlite3
from dataclasses import replace
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path

import pytest

from statewake.workspace import StateWakeWorkspace, WorkspaceRecordQuery
from statewake.workspace.projections import NormalizedRecordProjection
from statewake.workspace.xlsx_export import (
    XLSX_METADATA_SHEET,
    XLSX_RECORD_COLUMNS,
    XLSX_RECORD_SHEET,
    XlsxExportError,
    verify_xlsx,
    write_xlsx,
)


def _record(source_ref: str = "source-1") -> NormalizedRecordProjection:
    return NormalizedRecordProjection(
        record_id="record-1",
        receipt_id="receipt-1",
        artifact_digest="a" * 64,
        artifact_size=4,
        producer_id="producer-1",
        producer_type="test",
        producer_version=None,
        source_ref=source_ref,
        source_event_id=None,
        run_id=None,
        captured_at=datetime(2026, 1, 1, 12, tzinfo=UTC),
        sensitivity="internal",
        policy_id="default",
        verification_status=None,
        reliability_state=None,
        created_at=datetime(2026, 1, 1, 12, tzinfo=UTC),
    )


def test_xlsx_write_and_reopen_round_trip(tmp_path: Path) -> None:
    result = write_xlsx(
        (_record(),),
        tmp_path / "records.xlsx",
        schema_version="1",
        disclosure_max_sensitivity="internal",
        query_definition={"limit": 1},
    )
    verify_xlsx(result.output_path, expected_rows=1, expected_schema_version="1")

    from openpyxl import load_workbook

    workbook = load_workbook(result.output_path, read_only=True, data_only=False)
    assert tuple(workbook.sheetnames) == (XLSX_RECORD_SHEET, XLSX_METADATA_SHEET)
    sheet = workbook[XLSX_RECORD_SHEET]
    assert tuple(next(sheet.iter_rows(values_only=True))) == XLSX_RECORD_COLUMNS
    row = tuple(next(sheet.iter_rows(min_row=2, values_only=True)))
    assert row[0] == "record-1"
    assert row[6] is None
    assert row[8] is None
    assert row[10] == datetime(2026, 1, 1, 12, tzinfo=UTC)
    workbook.close()


def test_xlsx_preserves_formula_like_source_text_as_text(tmp_path: Path) -> None:
    result = write_xlsx(
        (_record(source_ref="=1+1"),),
        tmp_path / "records.xlsx",
        schema_version="1",
        disclosure_max_sensitivity="internal",
        query_definition={},
    )

    from openpyxl import load_workbook

    workbook = load_workbook(result.output_path, read_only=False, data_only=False)
    cell = workbook[XLSX_RECORD_SHEET]["H2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"
    workbook.close()


def test_xlsx_rejects_wrong_sheet_schema(tmp_path: Path) -> None:
    result = write_xlsx(
        (_record(),),
        tmp_path / "records.xlsx",
        schema_version="1",
        disclosure_max_sensitivity="internal",
        query_definition={},
    )

    from openpyxl import load_workbook

    workbook = load_workbook(result.output_path)
    workbook.remove(workbook[XLSX_METADATA_SHEET])
    workbook.save(result.output_path)
    workbook.close()

    with pytest.raises(XlsxExportError, match="sheet mismatch"):
        verify_xlsx(result.output_path, expected_rows=1, expected_schema_version="1")


def test_workspace_xlsx_export_matches_selected_projection_and_records_history(
    tmp_path: Path,
) -> None:
    with StateWakeWorkspace.open(tmp_path / "workspace") as workspace:
        workspace.ingest(
            b"one",
            producer_type="test",
            producer_id="producer",
            source_ref="source-1",
            captured_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
        workspace.ingest(
            b"two",
            producer_type="test",
            producer_id="producer",
            source_ref="source-2",
            captured_at=datetime(2026, 1, 2, tzinfo=UTC),
        )
        query = WorkspaceRecordQuery(limit=1)
        projection = workspace.project(query)
        result = workspace.export_xlsx(query, tmp_path / "records.xlsx")
        database_path = workspace.configuration.database_path

    with sqlite3.connect(database_path) as database:
        history = database.execute(
            "SELECT format, row_count, source_schema_version FROM exports "
            "WHERE export_id = ?",
            (result.export_id,),
        ).fetchone()

    from openpyxl import load_workbook

    workbook = load_workbook(result.output_path, read_only=True, data_only=False)
    row = tuple(
        next(workbook[XLSX_RECORD_SHEET].iter_rows(min_row=2, values_only=True))
    )
    assert result.row_count == len(projection.records) == 1
    assert row[0] == projection.records[0].record_id
    assert row[10] == projection.records[0].captured_at.replace(tzinfo=None)
    assert history == ("xlsx", 1, "1")
    workbook.close()


def test_xlsx_normalizes_aware_datetimes_to_utc(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    record = replace(
        _record(),
        captured_at=datetime(
            2026, 1, 1, 17, 30, tzinfo=timezone(timedelta(hours=5, minutes=30))
        ),
        created_at=datetime(
            2026, 1, 1, 17, 30, tzinfo=timezone(timedelta(hours=5, minutes=30))
        ),
    )
    result = write_xlsx(
        (record,),
        tmp_path / "records.xlsx",
        schema_version="1",
        disclosure_max_sensitivity="internal",
        query_definition={},
    )
    workbook = load_workbook(result.output_path, read_only=True, data_only=False)
    row = tuple(
        next(workbook[XLSX_RECORD_SHEET].iter_rows(min_row=2, values_only=True))
    )
    assert row[10] == datetime(2026, 1, 1, 12, tzinfo=UTC)
    assert row[15] == datetime(2026, 1, 1, 12, tzinfo=UTC)
    workbook.close()


def test_workspace_xlsx_export_applies_disclosure_ceiling(tmp_path: Path) -> None:
    with StateWakeWorkspace.open(tmp_path / "workspace") as workspace:
        workspace.ingest(
            b"one",
            producer_type="test",
            producer_id="producer",
            source_ref="source-1",
            captured_at=datetime(2026, 1, 1, tzinfo=UTC),
        )
        result = workspace.export_xlsx(
            WorkspaceRecordQuery(max_sensitivity="internal"),
            tmp_path / "records.xlsx",
            max_sensitivity="internal",
        )

    assert result.disclosure_max_sensitivity == "internal"
